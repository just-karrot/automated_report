import os
import io
import base64
import json
import concurrent.futures
import time
from flask import Flask, render_template, request, jsonify, Response, send_file
from dotenv import load_dotenv
from groq import Groq

load_dotenv()

app = Flask(__name__)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate_full_report():
    data = request.json
    api_key = data.get('api_key') or os.getenv('GROQ_API_KEY')
    description = data.get('description')
    event_category = data.get('event_category', 'technical')
    report_format = data.get('report_format', 'event_report')

    if not api_key:
        return jsonify({'error': 'API Key is required'}), 400
    
    if not description:
        return jsonify({'error': 'Event description is required'}), 400

    try:
        client = Groq(api_key=api_key)
        existing_data = data.get('existing_data')
        
        # --- PROJECT REPORT (Multi-Pass Generation) ---
        if report_format == 'project_report':
            # Determine if we are retrying failed sections
            is_retry = False
            if existing_data:
                failed_chapters = [i for i in range(1, 7) if str(existing_data.get(f'chapter_{i}_content', '')).startswith("Content generation failed")]
                if failed_chapters:
                    is_retry = True
                    result = existing_data
            
            if not is_retry:
                # Phase 1: Metadata, Titles, Abstract, References
                system_prompt_base = (
                    "You are an expert academic project coordinator. "
                    "Define the metadata, detailed abstract, references, and 6 relevant technical chapter titles for a dissertation. "
                    "Tone: Formal and Academic. Format: JSON schema provided. "
                    "You may use simple HTML for the abstract: <p> for paragraphs, <b> for bold, <i> for italics, and <ul>/<li> for lists if needed."
                )
                
                base_schema = {
                    "name": "project_report_base_schema",
                    "strict": False,
                    "schema": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "project_type": {"type": "string"},
                            "branch": {"type": "string"},
                            "dep_name": {"type": "string", "default": "Advance Computing"},
                            "session": {"type": "string"},
                            "professor_name": {"type": "string"},
                            "professor_designation": {"type": "string"},
                            "certificate_students_text": {"type": "string"},
                            "students": {
                                "type": "array",
                                "items": {"type": "object", "properties": {"name": {"type": "string"}, "reg_no": {"type": "string"}, "roll_no": {"type": "string"}}}
                            },
                            "abstract_title": {"type": "string", "default": "ABSTRACT"},
                            "abstract_content": {"type": "string"},
                            "keywords": {"type": "string"},
                            "chapter_1_title": {"type": "string"},
                            "chapter_2_title": {"type": "string"},
                            "chapter_3_title": {"type": "string"},
                            "chapter_4_title": {"type": "string"},
                            "chapter_5_title": {"type": "string"},
                            "chapter_6_title": {"type": "string"},
                            "references": {"type": "array", "items": {"type": "string"}}
                        },
                        "required": ["title", "chapter_1_title"]
                    }
                }

                prompt_base = f"Create the foundation for an engineering dissertation based on this description: {description}. Provide descriptive metadata, an extensive multi-paragraph abstract, 5-8 IEEE references, and 6 technical chapter titles (starting with CHAPTER X: ...)."
                
                response_base = client.chat.completions.create(
                    model="openai/gpt-oss-120b",
                    messages=[{"role": "system", "content": system_prompt_base}, {"role": "user", "content": prompt_base}],
                    response_format={"type": "json_schema", "json_schema": base_schema}
                )
                result = json.loads(response_base.choices[0].message.content or "{}")

            # Phase 2: Chapter Content Generation (Parallel)
            def generate_chapter_content(idx, chapter_title):
                chap_system_prompt = (
                    "You are a senior engineering student writing a highly detailed, professional project dissertation. "
                    "Write extensive, multi-paragraph content for the requested chapter. "
                    "Focus on technical depth, clear explanations, and academic rigor. "
                    "You are encouraged to use simple HTML to structure your content: "
                    "Use <p> for paragraphs, <b> for bold, <i> for italics, <ul>/<li> for bulleted lists, "
                    "and <table>/<tr>/<td> for data or comparisons. No raw markdown. "
                    "Do NOT use multiple <br> tags to simulate page breaks; the system will handle pagination. "
                    "The report must feel long and comprehensive (equivalent to 3-4 pages per chapter)."
                )
                chap_prompt = f"Write the complete, extremely detailed, multi-paragraph content for {chapter_title} of the project: {result.get('title')}. Context: {description}"
                
                # Internal retry loop (3 attempts)
                for attempt in range(3):
                    try:
                        chap_resp = client.chat.completions.create(
                            model="openai/gpt-oss-120b",
                            messages=[{"role": "system", "content": chap_system_prompt}, {"role": "user", "content": chap_prompt}],
                            response_format={
                                "type": "json_schema",
                                "json_schema": {
                                    "name": "chapter_content",
                                    "strict": False,
                                    "schema": {
                                        "type": "object",
                                        "properties": {"content": {"type": "string"}},
                                        "required": ["content"]
                                    }
                                }
                            }
                        )
                        return json.loads(chap_resp.choices[0].message.content or "{}").get("content", "")
                    except Exception as e:
                        if attempt < 2: # Don't sleep on last attempt
                            time.sleep(2 * (attempt + 1))
                        last_error = str(e)
                
                raise Exception(last_error)

            # Execute parallel calls (Reduced to 3 workers to avoid rate limits)
            with concurrent.futures.ThreadPoolExecutor(max_workers=3) as executor:
                future_to_idx = {}
                for i in range(1, 7):
                    title = result.get(f'chapter_{i}_title', f'CHAPTER {i}: CONTENT')
                    # ONLY generate if it's not a retry OR if it failed previously
                    content = str(result.get(f'chapter_{i}_content', ''))
                    if not content or content.startswith("Content generation failed"):
                        future_to_idx[executor.submit(generate_chapter_content, i, title)] = i
                
                for future in concurrent.futures.as_completed(future_to_idx):
                    idx = future_to_idx[future]
                    try:
                        result[f'chapter_{idx}_content'] = future.result()
                    except Exception as e:
                        result[f'chapter_{idx}_content'] = f"Content generation failed for this section. Error: {str(e)}"
            
            return jsonify(result)

        # --- EVENT REPORT (Single Pass, unchanged) ---
        else:
            system_prompts = {
                "cultural": "You are a professional report writer for a college cultural event. Write in a warm, lighthearted tone. MAINTAIN FORMALITY. JSON schema provided.",
                "sports": "You are a professional report writer for a college sports event. Focus on format, rules, and spirit. JSON schema provided.",
                "technical": "You are a professional report writer for a college technical event. Use precise terminology and emphasize innovation. JSON schema provided."
            }
            system_prompt = system_prompts.get(event_category, system_prompts["technical"])
            
            json_schema = {
                "name": "full_event_report_schema",
                "strict": True,
                "schema": {
                    "type": "object",
                    "properties": {
                        "event_name": {"type": "string"},
                        "event_date": {"type": "string"},
                        "event_duration": {"type": "string"},
                        "event_type": {"type": "string"},
                        "organized_by": {"type": "string"},
                        "coordinators": {"type": "string"},
                        "learning_outcome_1": {"type": "string"},
                        "learning_outcome_2": {"type": "string"},
                        "learning_outcome_3": {"type": "string"},
                        "learning_outcome_4": {"type": "string"},
                        "introduction": {"type": "string"},
                        "details_of_the_event": {"type": "string"},
                        "description_of_the_event": {"type": "string"},
                        "beneficiaries": {"type": "string"},
                        "session_overview": {"type": "string"},
                        "key_highlights": {"type": "array", "items": {"type": "string"}},
                        "conclusion": {"type": "string"}
                    },
                    "required": [
                        "event_name", "event_date", "event_duration", "event_type", 
                        "organized_by", "coordinators", "learning_outcome_1", 
                        "learning_outcome_2", "learning_outcome_3", "learning_outcome_4", 
                        "introduction", "details_of_the_event", "description_of_the_event", 
                        "beneficiaries", "session_overview", "key_highlights", "conclusion"
                    ],
                    "additionalProperties": False
                }
            }
            
            prompt = f"Extract and professionally expand the following event description into a full formal event report JSON payload: {description}"
            response = client.chat.completions.create(
                model="openai/gpt-oss-120b",
                messages=[{"role": "system", "content": system_prompt}, {"role": "user", "content": prompt}],
                response_format={"type": "json_schema", "json_schema": json_schema}
            )
            return jsonify(json.loads(response.choices[0].message.content or "{}"))

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/render', methods=['POST'])
def render_report_preview():
    data = request.json
    report_format = data.get('report_format', 'event_report')
    template = 'project_schema.html' if report_format == 'project_report' else 'schema.html'
    return render_template(template, **data)

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from lxml import html

def add_html_to_docx(html_content, doc_or_cell):
    """Parses HTML and adds corresponding elements to a python-docx document or table cell."""
    if not html_content:
        return
    
    try:
        # Wrap in a div to ensure a single root for lxml
        tree = html.fromstring(f"<div>{html_content}</div>")
    except Exception:
        doc_or_cell.add_paragraph(html_content)
        return

    def process_element(element, parent_docx):
        # Handle child nodes
        for child in element.xpath('node()'):
            if isinstance(child, html.HtmlElement):
                tag = child.tag.lower()
                
                if tag == 'p':
                    p = parent_docx.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    process_element(child, p)
                
                elif tag == 'br':
                    # If parent is a paragraph, add a break
                    if hasattr(parent_docx, 'add_run'):
                        parent_docx.add_run().add_break()
                    else:
                        parent_docx.add_paragraph()
                
                elif tag in ['b', 'strong']:
                    if hasattr(parent_docx, 'add_run'):
                        run = parent_docx.add_run()
                        run.bold = True
                        # We need to process children of <b> to handle nested tags like <b><i>text</i></b>
                        # But for simplicity, if it's just text:
                        run.text = child.text if child.text else ""
                        process_element(child, parent_docx)
                    else:
                        p = parent_docx.add_paragraph()
                        run = p.add_run()
                        run.bold = True
                        run.text = child.text if child.text else ""
                        process_element(child, p)

                elif tag in ['i', 'em']:
                    if hasattr(parent_docx, 'add_run'):
                        run = parent_docx.add_run()
                        run.italic = True
                        run.text = child.text if child.text else ""
                        process_element(child, parent_docx)
                    else:
                        p = parent_docx.add_paragraph()
                        run = p.add_run()
                        run.italic = True
                        run.text = child.text if child.text else ""
                        process_element(child, p)

                elif tag == 'ul':
                    for li in child.xpath('./li'):
                        p = parent_docx.add_paragraph(style='List Bullet')
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        process_element(li, p)
                
                elif tag == 'ol':
                    for li in child.xpath('./li'):
                        p = parent_docx.add_paragraph(style='List Number')
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        process_element(li, p)

                elif tag == 'table':
                    rows = child.xpath('.//tr')
                    if rows:
                        # Find max columns
                        max_cols = 0
                        for row in rows:
                            max_cols = max(max_cols, len(row.xpath('./td | ./th')))
                        
                        if max_cols > 0:
                            table = parent_docx.add_table(rows=len(rows), cols=max_cols)
                            table.style = 'Table Grid'
                            for r_idx, row in enumerate(rows):
                                for c_idx, cell in enumerate(row.xpath('./td | ./th')):
                                    if c_idx < max_cols:
                                        process_element(cell, table.cell(r_idx, c_idx))

                elif tag in ['td', 'th', 'li']:
                    # These are containers, just process their children
                    process_element(child, parent_docx)
                
                else:
                    # Unknown tag, just process children
                    process_element(child, parent_docx)
                    
            elif isinstance(child, str):
                # This is text content
                text = str(child).strip()
                if text:
                    if hasattr(parent_docx, 'add_run'):
                        # If the current parent is a paragraph, add text to it
                        parent_docx.add_run(text)
                    elif hasattr(parent_docx, 'add_paragraph'):
                        # If the current parent is a Document or Cell, create a paragraph
                        p = parent_docx.add_paragraph(text)
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    process_element(tree, doc_or_cell)

@app.route('/export-docx', methods=['POST'])
def export_docx():
    data = request.json
    report_format = data.get('report_format', 'event_report')
    doc = Document()
    
    # Global Style Settings
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # Heading Styles to Black
    for style_name in ['Heading 1', 'Heading 2']:
        h_style = doc.styles[style_name]
        h_font = h_style.font
        h_font.name = 'Times New Roman'
        h_font.color.rgb = RGBColor(0, 0, 0)
        h_font.bold = True
        if style_name == 'Heading 1':
            h_font.size = Pt(16)
        else:
            h_font.size = Pt(14)

    if report_format == 'project_report':
        report_name = data.get('title', 'Project_Report')
        
        # Configure Footer for all pages except the first
        section = doc.sections[0]
        section.different_first_page_header_footer = True
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.text = f"Department of {data.get('dep_name', 'Advance Computing')}, Poornima College of Engineering"
        footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # --- COVER PAGE ---
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run(data.get('title', '[TITLE]').upper())
        run.bold = True
        run.font.size = Pt(16)
        
        type_p = doc.add_paragraph()
        type_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        type_p.add_run(f"A {data.get('project_type', '[Type]')} Report submitted in partial fulfilment of the requirements of")
        
        degree_p = doc.add_paragraph()
        degree_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        degree_p.add_run("The award of the degree of")
        
        btech_p = doc.add_paragraph()
        btech_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = btech_p.add_run("Bachelor of Technology")
        run.bold = True
        run.font.size = Pt(14)
        
        in_p = doc.add_paragraph()
        in_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        in_p.add_run("in")
        
        branch_p = doc.add_paragraph()
        branch_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = branch_p.add_run(data.get('branch', '[Branch]'))
        run.bold = True
        run.font.size = Pt(14)
        
        by_p = doc.add_paragraph()
        by_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        by_p.add_run("by")
        
        for student in data.get('students', []):
            sp = doc.add_paragraph()
            sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = sp.add_run(f"{student.get('name')}, Reg No: {student.get('reg_no')}")
            run.bold = True
            
        guide_p = doc.add_paragraph("Under the guidance of")
        guide_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        prof_p = doc.add_paragraph()
        prof_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = prof_p.add_run(data.get('professor_name', '[Professor Name]'))
        run.bold = True
        
        doc.add_paragraph(data.get('professor_designation', 'Assistant Professor')).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Department of Advance Computing").alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Logo
        logo_path = os.path.join(app.static_folder, 'images', 'pce_logo.jpg')
        if os.path.exists(logo_path):
            doc.add_picture(logo_path, width=Inches(1.2))
            last_p = doc.paragraphs[-1]
            last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        doc.add_paragraph(f"(Session {data.get('session', '2025-26')})").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Department of Advance Computing").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Poornima College of Engineering").alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("ISI-6, RIICO Institutional Area, Sitapura, Jaipur – 302022").alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break() # Break after Cover Page
        
        # --- CERTIFICATES ---
        doc.add_heading('DEPARTMENT CERTIFICATE', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        cert_p = doc.add_paragraph("This is to certify that ")
        cert_p.add_run(data.get('certificate_students_text', '[Students]')).bold = True
        cert_p.add_run(" of the IV semester Department of Advance Computing, has submitted this Project report entitled ")
        cert_p.add_run(data.get('title', '[Title]')).bold = True
        cert_p.add_run(f" under the supervision of {data.get('professor_name')}, {data.get('professor_designation')}, Department of Advance Computing, working in division of Advance Computing as per the requirements of the Bachelor of Technology program at Poornima College of Engineering, Jaipur affiliated by Rajasthan Technical University.")
        
        doc.add_paragraph()
        doc.add_paragraph("Dr. Amol Saxena").bold = True
        doc.add_paragraph("Head, Department of Advanced Computing")
        
        doc.add_page_break() # Break after Dept Certificate
        
        doc.add_heading("CANDIDATE'S DECLARATION", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        decl_p = doc.add_paragraph("We hereby declare that the work which is being presented in this project report entitled ")
        decl_p.add_run(data.get('title', '[Title]')).bold = True
        decl_p.add_run(" in the partial fulfilment for the award of the Degree of Bachelor of Technology in CSE(Cyber Security), submitted in the Department of Advanced Computing, Poornima College of Engineering, Jaipur, is an authentic record of our work done during the period from ")
        decl_p.add_run(data.get('session', '[Session]')).bold = True
        decl_p.add_run(f" under the supervision and guidance of {data.get('professor_name')}, {data.get('professor_designation')}, Department of Advanced Computing.")
        doc.add_paragraph("We have not submitted the matter embodied in this project report for the award of any other degree.")
        
        doc.add_paragraph()
        for student in data.get('students', []):
            doc.add_paragraph(f"Name of Candidate: {student.get('name')}")
            doc.add_paragraph(f"Registration no: {student.get('reg_no')}")
            doc.add_paragraph()

        doc.add_page_break() # Break after Candidate Declaration
        
        doc.add_heading("SUPERVISOR'S CERTIFICATE", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("This is to certify that, to the best of my knowledge, the candidate's above statement is correct.")
        doc.add_paragraph()
        doc.add_paragraph(data.get('professor_name', '[Professor Name]')).bold = True
        doc.add_paragraph(data.get('professor_designation', '[Designation]'))
        doc.add_paragraph("Department of Advanced Computing")
        
        doc.add_page_break() # Break after Supervisor Certificate
        
        doc.add_heading("ACKNOWLEDGEMENT", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        ack_p = doc.add_paragraph("We would like to convey our profound sense of reverence and admiration to my supervisor, ")
        ack_p.add_run(f"{data.get('professor_name')}, {data.get('professor_designation')} in the Department of Advance Computing at Poornima College of Engineering").bold = True
        ack_p.add_run(", for their intense concern, attention, priceless direction, guidance, and encouragement throughout this research work.")
        
        doc.add_paragraph("We are grateful to Dr. Mahesh Bundele, Principal & Director, and Dr. Pankaj Dhemla, Vice-Principal of Poornima College of Engineering, for providing the necessary resources and a conducive environment to carry out this project.")
        doc.add_paragraph("Our special heartfelt gratitude goes to Dr. Amol Saxena, HOD, and Dr. Kamlesh Gautam, Dy. HOD, Department of Advanced Computing, for unvarying support, guidance, and motivation during this project work.")
        doc.add_paragraph("We would like to express our deep sense of gratitude towards the management of Poornima College of Engineering, including Shri Shashikant Singhi, Chairman, Poornima Group, Mr. M. K. M. Shah, Director General, Poornima Group, and Ar. Rahul Singhi, Director of Poornima Group, for providing all the necessary resources and facilities required to complete this project.")
        doc.add_paragraph("We would like to take the opportunity to express our thanks to all faculty members of the Department for their kind support, technical guidance, and inspiration throughout the course.")
        doc.add_paragraph("We are also thankful to the non-teaching staff of the department for their support in the preparation of this dissertation work.")
        doc.add_paragraph("We are deeply thankful to my parents and all other family members for their blessings and inspiration. Last, but not least, we would like to give special thanks to God who enabled me to complete my dissertation on time.")
        
        doc.add_paragraph()
        for student in data.get('students', []):
            doc.add_paragraph(f"{student.get('name')}, Department of Advanced Computing, {student.get('roll_no')}").bold = True

        doc.add_page_break() # Break after Acknowledgement
        
        doc.add_heading("TABLE OF CONTENTS", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_table = doc.add_table(rows=1, cols=2)
        toc_table.cell(0, 0).text = "Contents"
        toc_table.cell(0, 1).text = "Page No."
        toc_table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        toc_items = [
            (data.get('abstract_title', 'ABSTRACT'), "1"),
            (f"CHAPTER 1: {data.get('chapter_1_title')}", "2"),
            (f"CHAPTER 2: {data.get('chapter_2_title')}", "3"),
            (f"CHAPTER 3: {data.get('chapter_3_title')}", "4"),
            (f"CHAPTER 4: {data.get('chapter_4_title')}", "5"),
            (f"CHAPTER 5: {data.get('chapter_5_title')}", "6"),
            (f"CHAPTER 6: {data.get('chapter_6_title')}", "7"),
            ("REFERENCES", "8")
        ]
        for item, page in toc_items:
            row = toc_table.add_row().cells
            row[0].text = item
            row[1].text = page
            row[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            
        doc.add_page_break()
        
        doc.add_heading("LIST OF FIGURES", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        fig_table = doc.add_table(rows=1, cols=4)
        fig_table.style = 'Table Grid'
        hdr = fig_table.rows[0].cells
        hdr[0].text = "S. No."
        hdr[1].text = "Fig. No."
        hdr[2].text = "Description"
        hdr[3].text = "Page No."
        for i in range(2): fig_table.add_row() # Placeholder rows

        doc.add_page_break()

        doc.add_heading("LIST OF ACRONYMS", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        acr_table = doc.add_table(rows=1, cols=3)
        acr_table.style = 'Table Grid'
        hdr = acr_table.rows[0].cells
        hdr[0].text = "Serial Number"
        hdr[1].text = "ACRONYM"
        hdr[2].text = "FULL FORM"
        for i in range(2): acr_table.add_row() # Placeholder rows
        
        doc.add_page_break()
        
        # --- ABSTRACT ---
        doc.add_heading(data.get('abstract_title', 'ABSTRACT'), level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_html_to_docx(data.get('abstract_content', ''), doc)
        doc.add_paragraph(f"Keywords: {data.get('keywords', '')}")
        
        doc.add_page_break() # Break after Keywords (before chapters)
        
        # --- CHAPTERS ---
        for i in range(1, 7):
            title = data.get(f'chapter_{i}_title')
            content = data.get(f'chapter_{i}_content', '')
            if title:
                doc.add_heading(title, level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_html_to_docx(content, doc)
                doc.add_page_break()
                
        # --- REFERENCES ---
        doc.add_heading('REFERENCES', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ref in data.get('references', []):
            doc.add_paragraph(ref, style='List Bullet')

    else:
        # --- EVENT REPORT ---
        report_name = data.get('event_name', 'Event_Report')
        
        # Header Image
        header_path = os.path.join(app.static_folder, 'images', 'header.jpg')
        if os.path.exists(header_path):
            doc.add_picture(header_path, width=Inches(6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph("A report on").alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_p.add_run(data.get('event_name', '').upper())
        run.bold = True
        run.underline = True
        
        details = [
            ("NAME OF THE EVENT:", data.get('event_name')),
            ("DATE AND DURATION:", f"{data.get('event_date')} - {data.get('event_duration')}"),
            ("TYPE OF ACTIVITY:", data.get('event_type')),
            ("ORGANIZED BY:", data.get('organized_by')),
            ("COORDINATORS:", data.get('coordinators', '').replace('<br>', '\n'))
        ]
        
        for label, val in details:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(label)
            run.bold = True
            p.add_run(f" {val}")

        doc.add_page_break()
        
        # Outcomes, Content, etc.
        doc.add_heading('LEARNING OUTCOMES', level=2)
        for i in range(1, 5):
            doc.add_paragraph(f"CO{i}: {data.get(f'learning_outcome_{i}', '')}")
            
        sections = [
            ("INTRODUCTION", data.get('introduction')),
            ("DETAILS OF THE EVENT", data.get('details_of_the_event')),
            ("DESCRIPTION OF THE EVENT", data.get('description_of_the_event')),
            ("BENEFICIARIES", data.get('beneficiaries')),
            ("SESSION OVERVIEW", data.get('session_overview')),
            ("CONCLUSION", data.get('conclusion'))
        ]
        
        for title, content in sections:
            doc.add_heading(title, level=2)
            add_html_to_docx(content, doc)

    # Save to buffer
    f = io.BytesIO()
    doc.save(f)
    f.seek(0)

    # Clean filename
    safe_name = "".join(c for c in report_name if c.isalnum() or c in (' ', '-', '_')).strip()
    filename = f"{safe_name or 'Report'}.docx"

    return send_file(
        f,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

if __name__ == '__main__':
    app.run(debug=True)

